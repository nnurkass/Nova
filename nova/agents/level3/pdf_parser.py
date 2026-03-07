"""Document parser contracts for PDF/DOCX tender attachments."""

from __future__ import annotations

import asyncio
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from langchain_core.tools import StructuredTool
from pydantic import BaseModel, Field
from PyPDF2 import PdfReader

from nova.integrations.abc.models import ABCWork


SECTION_KEYS = (
    "technical_specification",
    "requirements",
    "work_scope",
)

SECTION_HEADINGS: dict[str, tuple[str, ...]] = {
    "technical_specification": (
        "technical specification",
        "technical assignment",
        "техническое задание",
        "техническая спецификация",
        "тех задание",
        "тз",
    ),
    "requirements": (
        "requirements",
        "qualification requirements",
        "требования",
        "квалификационные требования",
    ),
    "work_scope": (
        "work scope",
        "scope of work",
        "volumes of work",
        "объемы работ",
        "перечень работ",
    ),
}

WORK_NAME_HINTS = (
    "work",
    "works",
    "construction",
    "install",
    "concrete",
    "pouring",
    "earth",
    "монтаж",
    "работ",
    "устройство",
    "кладка",
    "землян",
    "бетон",
)

UNIT_ALIASES = {
    "м2": "m2",
    "м²": "m2",
    "кв.м": "m2",
    "м3": "m3",
    "м³": "m3",
    "куб.м": "m3",
    "шт.": "pcs",
    "шт": "pcs",
    "ед.": "pcs",
    "ед": "pcs",
    "тонна": "t",
    "тонн": "t",
    "тн": "t",
}

LINE_ITEM_PATTERN = re.compile(
    r"^(?:(?P<code>[A-Za-zА-Яа-я0-9_.-]{1,24})\s*[-:|;]\s*)?"
    r"(?P<name>[A-Za-zА-Яа-я0-9().,%/\-+ ]{4,}?)\s+"
    r"(?P<quantity>\d+(?:[.,]\d+)?)\s*"
    r"(?P<unit>[A-Za-zА-Яа-я0-9./²³%-]{1,12})$"
)


class DocumentParseError(ValueError):
    """Raised when a document cannot be parsed into a structured payload."""


class WorkExtractionToolInput(BaseModel):
    """Input schema for the work extraction tool."""

    file_path: str = Field(..., min_length=1, description="Path to a local PDF or DOCX file.")


def _normalize_text(text: str) -> str:
    """Normalize spacing while preserving paragraph boundaries."""
    lines = [" ".join(line.strip().split()) for line in text.replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def _empty_sections() -> dict[str, str]:
    return {key: "" for key in SECTION_KEYS}


def _normalize_heading(value: str) -> str:
    lowered = value.lower().replace("ё", "е")
    cleaned = re.sub(r"[^a-zа-я0-9 ]+", " ", lowered, flags=re.IGNORECASE)
    return " ".join(cleaned.split())


def _detect_section_heading(line: str) -> str | None:
    normalized = _normalize_heading(line)
    if not normalized:
        return None

    for key, aliases in SECTION_HEADINGS.items():
        for alias in aliases:
            if normalized == alias:
                return key
            if normalized.startswith(f"{alias} ") and len(normalized) <= len(alias) + 4:
                return key
    return None


def _split_sections(text: str) -> dict[str, str]:
    """Best-effort section splitter based on heading aliases."""

    normalized = _normalize_text(text)
    sections = _empty_sections()
    if not normalized:
        return sections

    active_key = "technical_specification"
    bucket: dict[str, list[str]] = {key: [] for key in SECTION_KEYS}

    for line in normalized.splitlines():
        heading_key = _detect_section_heading(line)
        if heading_key is not None:
            active_key = heading_key
            continue
        bucket[active_key].append(line)

    for key in SECTION_KEYS:
        sections[key] = "\n".join(bucket[key]).strip()

    if not any(sections.values()):
        sections["technical_specification"] = normalized

    return sections


def _normalize_table_rows(rows: list[list[str]]) -> dict[str, list[list[str]] | list[str]]:
    if not rows:
        return {"headers": [], "rows": []}

    width = max(len(row) for row in rows)
    normalized_rows: list[list[str]] = []
    for row in rows:
        normalized = row + [""] * (width - len(row))
        if any(cell.strip() for cell in normalized):
            normalized_rows.append([_normalize_text(cell) for cell in normalized])

    if not normalized_rows:
        return {"headers": [], "rows": []}

    return {
        "headers": normalized_rows[0],
        "rows": normalized_rows[1:],
    }


def _extract_docx_with_python_docx(path: Path) -> tuple[list[str], list[dict[str, list[list[str]] | list[str]]]]:
    from docx import Document  # type: ignore[import-not-found]

    document = Document(str(path))
    paragraphs = [_normalize_text(paragraph.text) for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[dict[str, list[list[str]] | list[str]]] = []

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [_normalize_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        normalized = _normalize_table_rows(rows)
        if normalized["headers"]:
            tables.append(normalized)

    return paragraphs, tables


def _extract_docx_with_xml(path: Path) -> tuple[list[str], list[dict[str, list[list[str]] | list[str]]]]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        with zipfile.ZipFile(path, "r") as archive:
            document_xml = archive.read("word/document.xml")
    except Exception as exc:  # pragma: no cover - defensive path
        raise DocumentParseError(f"Failed to read DOCX zip container: {path}") from exc

    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:  # pragma: no cover - defensive path
        raise DocumentParseError(f"Invalid DOCX XML payload: {path}") from exc

    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:body/w:p", ns):
        fragments = [node.text or "" for node in paragraph.findall(".//w:t", ns)]
        text = _normalize_text(" ".join(fragments))
        if text:
            paragraphs.append(text)

    tables: list[dict[str, list[list[str]] | list[str]]] = []
    for table in root.findall(".//w:body/w:tbl", ns):
        rows: list[list[str]] = []
        for row in table.findall("./w:tr", ns):
            cells: list[str] = []
            for cell in row.findall("./w:tc", ns):
                fragments = [node.text or "" for node in cell.findall(".//w:t", ns)]
                cells.append(_normalize_text(" ".join(fragments)))
            if any(cells):
                rows.append(cells)
        normalized = _normalize_table_rows(rows)
        if normalized["headers"]:
            tables.append(normalized)

    return paragraphs, tables


def _validate_path(file_path: str | Path, *, expected_suffix: str) -> Path:
    path = Path(file_path).expanduser()
    if not path.exists():
        raise DocumentParseError(f"File not found: {path}")
    if path.suffix.lower() != expected_suffix:
        raise DocumentParseError(
            f"Unsupported file extension for parser {expected_suffix}: {path.suffix.lower() or '<none>'}"
        )
    return path


def parse_pdf(file_path: str | Path) -> dict[str, Any]:
    """Parse PDF content into a unified intermediate document structure."""

    path = _validate_path(file_path, expected_suffix=".pdf")
    try:
        reader = PdfReader(str(path))
        text_chunks = [page.extract_text() or "" for page in reader.pages]
        text = _normalize_text("\n".join(text_chunks))
    except Exception as exc:  # pragma: no cover - defensive, triggered in broken-file test
        raise DocumentParseError(f"Failed to parse PDF file: {path}") from exc

    if not text:
        raise DocumentParseError(f"No extractable text found in PDF: {path}")

    return {
        "document_type": "pdf",
        "file_path": str(path),
        "text": text,
        "sections": _split_sections(text),
        "tables": [],
        "metadata": {
            "page_count": len(reader.pages),
        },
    }


def parse_docx(file_path: str | Path) -> dict[str, Any]:
    """Parse DOCX content into a unified intermediate document structure."""

    path = _validate_path(file_path, expected_suffix=".docx")
    parser_backend = "python-docx"
    try:
        paragraphs, tables = _extract_docx_with_python_docx(path)
    except Exception:
        parser_backend = "xml-fallback"
        paragraphs, tables = _extract_docx_with_xml(path)

    text = _normalize_text("\n".join(paragraphs))
    if not text and not tables:
        raise DocumentParseError(f"No extractable content found in DOCX: {path}")

    return {
        "document_type": "docx",
        "file_path": str(path),
        "text": text,
        "sections": _split_sections(text),
        "tables": tables,
        "metadata": {
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "parser_backend": parser_backend,
        },
    }


def parse_document(file_path: str | Path) -> dict[str, Any]:
    """Dispatch document parsing by file extension."""

    path = Path(file_path).expanduser()
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise DocumentParseError(f"Unsupported document type: {suffix or '<none>'}")


def _json_dumps(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _tool_error(tool_name: str, file_path: str, exc: Exception) -> str:
    return _json_dumps(
        {
            "tool": tool_name,
            "status": "error",
            "file_path": file_path,
            "error_type": exc.__class__.__name__,
            "error": str(exc),
        }
    )


def _normalize_unit(unit: str) -> str:
    normalized = unit.strip().lower().replace(" ", "")
    return UNIT_ALIASES.get(normalized, normalized)


def _parse_number(raw: str) -> float:
    return float(raw.replace(" ", "").replace(",", "."))


def _looks_like_work_name(name: str) -> bool:
    normalized = _normalize_heading(name)
    return any(hint in normalized for hint in WORK_NAME_HINTS)


def _find_column_index(headers: list[str], aliases: tuple[str, ...]) -> int | None:
    normalized_headers = [_normalize_heading(header) for header in headers]
    for index, header in enumerate(normalized_headers):
        if any(alias in header for alias in aliases):
            return index
    return None


def _record_key(record: dict[str, Any]) -> tuple[str, str, float]:
    return (
        str(record["name"]).strip().lower(),
        str(record["unit"]).strip().lower(),
        float(record["quantity"]),
    )


def _extract_work_from_tables(tables: list[dict[str, list[list[str]] | list[str]]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    generated = 1

    for table in tables:
        headers = [str(value) for value in table.get("headers", [])]
        rows = table.get("rows", [])
        if not headers or not isinstance(rows, list):
            continue

        name_idx = _find_column_index(headers, ("name", "наименование", "работ", "work"))
        quantity_idx = _find_column_index(headers, ("qty", "quantity", "кол", "объем", "volume"))
        unit_idx = _find_column_index(headers, ("unit", "ед", "изм", "uom"))
        code_idx = _find_column_index(headers, ("code", "код", "шифр", "позиция"))
        if name_idx is None or quantity_idx is None or unit_idx is None:
            continue

        for row in rows:
            if not isinstance(row, list):
                continue
            if max(name_idx, quantity_idx, unit_idx) >= len(row):
                continue
            raw_name = str(row[name_idx]).strip()
            raw_quantity = str(row[quantity_idx]).strip()
            raw_unit = str(row[unit_idx]).strip()
            if not raw_name or not raw_quantity or not raw_unit:
                continue
            if not _looks_like_work_name(raw_name):
                continue

            try:
                quantity = _parse_number(raw_quantity)
            except ValueError:
                continue

            raw_code = ""
            if code_idx is not None and code_idx < len(row):
                raw_code = str(row[code_idx]).strip()
            code = raw_code or f"W-{generated:03d}"
            generated += 1

            record = ABCWork(
                code=code,
                name=raw_name,
                unit=_normalize_unit(raw_unit),
                quantity=quantity,
            ).model_dump(mode="json")
            records.append(record)

    return records


def _extract_work_from_text_section(text: str, *, code_start: int) -> tuple[list[dict[str, Any]], int]:
    records: list[dict[str, Any]] = []
    next_code = code_start
    for raw_line in text.splitlines():
        line = raw_line.strip(" -•\t")
        if not line:
            continue

        match = LINE_ITEM_PATTERN.match(line)
        if not match:
            continue

        name = match.group("name").strip()
        if not _looks_like_work_name(name):
            continue

        try:
            quantity = _parse_number(match.group("quantity"))
        except ValueError:
            continue

        code = (match.group("code") or "").strip() or f"W-{next_code:03d}"
        next_code += 1
        record = ABCWork(
            code=code,
            name=name,
            unit=_normalize_unit(match.group("unit")),
            quantity=quantity,
        ).model_dump(mode="json")
        records.append(record)

    return records, next_code


def _extract_work_records(document: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    records = _extract_work_from_tables(document.get("tables", []))
    source_sections: set[str] = set()

    next_code = len(records) + 1
    sections = document.get("sections", {})
    if isinstance(sections, dict):
        for section_name in ("work_scope", "technical_specification"):
            section_text = sections.get(section_name, "")
            if not isinstance(section_text, str) or not section_text:
                continue

            section_records, next_code = _extract_work_from_text_section(section_text, code_start=next_code)
            if section_records:
                source_sections.add(section_name)
                records.extend(section_records)

    unique_records: list[dict[str, Any]] = []
    seen: set[tuple[str, str, float]] = set()
    for record in records:
        key = _record_key(record)
        if key in seen:
            continue
        seen.add(key)
        unique_records.append(record)

    if document.get("tables"):
        source_sections.add("tables")

    return unique_records, sorted(source_sections)


def _extract_work_payload(file_path: str) -> dict[str, Any]:
    document = parse_document(file_path)
    work_list, source_sections = _extract_work_records(document)
    return {
        "tool": "extract_work_list",
        "status": "ok",
        "file_path": document["file_path"],
        "document_type": document["document_type"],
        "work_list": work_list,
        "source_sections": source_sections,
        "counts": {
            "work_items": len(work_list),
        },
    }


def _extract_work_list_sync(*, file_path: str) -> str:
    try:
        return _json_dumps(_extract_work_payload(file_path))
    except Exception as exc:  # pragma: no cover - validated by tool error tests
        return _tool_error("extract_work_list", file_path, exc)


async def _extract_work_list_async(*, file_path: str) -> str:
    return await asyncio.to_thread(_extract_work_list_sync, file_path=file_path)


extract_work_list = StructuredTool.from_function(
    func=_extract_work_list_sync,
    coroutine=_extract_work_list_async,
    name="extract_work_list",
    description=(
        "Extract a structured list of construction work items from a local PDF/DOCX file. "
        "Returns JSON with normalized code, name, unit, quantity, and source section metadata."
    ),
    args_schema=WorkExtractionToolInput,
)


__all__ = [
    "DocumentParseError",
    "SECTION_KEYS",
    "WorkExtractionToolInput",
    "extract_work_list",
    "parse_docx",
    "parse_document",
    "parse_pdf",
]
